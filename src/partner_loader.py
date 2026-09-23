"""Explicit adapters for the inspected partner workbooks; no invented defaults."""

from dataclasses import dataclass
from pathlib import Path
import re

import numpy as np
import pandas as pd
from openpyxl import load_workbook


FILES = {
    "IEK": {
        "sales": "Ежемесячные продажи в количественном выражении за последние 2 года.xlsx",
        "inventory": "Ежемесячные остатки продукции за последние 2 года  ИЭК.xlsx",
        "documents": "Динамика продаж_2025-2026.xlsx",
        "moq": "MOQ  ИЭК.xlsx", "transit": "Путь ИЭК 22.09.2026.xlsx",
        "seasonality": "Сезонность ИЭК.xlsx",
    },
    "Systeme Electric": {
        "sales": "Ежемесячные продажи в кол-м выражении SystemElectric 2024-2026.xlsx",
        "inventory": "Ежемесячные остатки SystemElectric 2024-2026.xlsx",
        "documents": "Динамика продаж_Syseme Electric_2025-2026.xlsx",
        "moq": "MOQ SystemElectric.xlsx",
        "transit": "Товар в пути_SystemElectric на 22.09.2026.xlsx",
        "seasonality": "Сезонность SystemElectric 2024-2026.xlsx",
    },
}
MONTHS = {name: i + 1 for i, name in enumerate(
    ["янв", "фев", "мар", "апр", "май", "июн", "июл", "авг", "сен", "окт", "ноя", "дек"])}


@dataclass
class PartnerData:
    catalog: pd.DataFrame
    history: pd.DataFrame
    documents: pd.DataFrame
    inbound: pd.DataFrame
    seasonality: dict
    warehouse_stock: pd.DataFrame
    inspection: list
    issues: list
    as_of: pd.Timestamp


def workbook_fingerprint(raw_dir):
    """Cache key changes on add/remove/replace/update of any supplied workbook."""
    return tuple((p.name, p.stat().st_size, p.stat().st_mtime_ns)
                 for p in sorted(Path(raw_dir).glob("*.xlsx")))


def month_header(value):
    text = str(value).strip().lower()
    year = re.search(r"\b(20\d{2})\b", text)
    month = MONTHS.get(text[:3])
    return pd.Timestamp(int(year[1]), month, 1) if year and month else None


def _table(rows, header, required, source):
    columns = [str(v).strip() if v is not None else f"_blank_{i}" for i, v in enumerate(rows[header])]
    missing = set(required) - set(columns)
    if missing:
        raise ValueError(f"{source}: unexpected workbook schema; missing {sorted(missing)}")
    return pd.DataFrame(rows[header + 1:], columns=columns)


def _key_rows(frame, key):
    result = frame.loc[frame[key].notna()].copy()
    # Preserve text identifiers (including leading zeroes); Excel numeric codes
    # may arrive as floats, but must join the same integer code in other sheets.
    result["sku"] = result[key].map(
        lambda value: str(int(value)) if isinstance(value, (int, float))
        and np.isfinite(value) and float(value).is_integer() else str(value).strip()).astype(str)
    # The IEK transit workbook contains an otherwise empty placeholder code 0.
    # It is not a product and must never become a selectable SKU.
    return result.loc[~result.sku.str.casefold().isin(
        ["", "0", "0.0", "none", "nan", "<na>", "итого"])]


def normalize_monthly(frame, key, name, value_name):
    """Unpivot Russian month columns; blanks remain unknown, totals are ignored."""
    frame = _key_rows(frame, key)
    columns = {c: month_header(c) for c in frame.columns if month_header(c) is not None}
    if not columns:
        raise ValueError("No recognized monthly columns")
    if frame.sku.duplicated().any():
        raise ValueError("Duplicate SKU in monthly report; aggregation is ambiguous")
    long = frame.melt(id_vars=["sku", name], value_vars=list(columns), var_name="month", value_name=value_name)
    long["date"] = long["month"].map(columns)
    long[value_name] = pd.to_numeric(long[value_name], errors="coerce")
    return long.rename(columns={name: "product_name"})[["sku", "product_name", "date", value_name]]


def infer_stockout_evidence(history):
    """A zero snapshot supports a POSSIBLE outage, never invented duration days."""
    result = history.copy()
    result["possible_stockout"] = pd.to_numeric(result["historical_stock"], errors="coerce").eq(0)
    result["stockout_evidence"] = np.where(result.possible_stockout, "zero_inventory_snapshot", "unavailable")
    result["stockout_days"] = np.nan
    return result


def detect_document_anomalies(documents):
    """High-confidence statistical one-off documents, not customer attribution.

    Aggregate signed shipment lines per SKU/document first. At least eight
    positive documents are required. A spike exceeds median + 6 robust sigmas,
    five times the median and 10 units. Comparable repeat large documents
    (>=80% of a candidate) prevent automatic exclusion. customer_id, if supplied,
    is retained as optional anonymized metadata, never inferred from document IDs.
    """
    result = documents.copy()
    keys = ["brand", "sku", "document_id", "document_type", "warehouse"]
    result["source_line_count"] = 1
    if result.duplicated(keys).any():
        # Preserve signed line quantities for audit while assessing the whole document.
        result["source_quantities"] = result.quantity.map(lambda q: [q])
        aggregations = {c: "first" for c in result.columns if c not in keys}
        aggregations.update(quantity=lambda q: q.sum(min_count=1), source_line_count="sum",
                            source_quantities=lambda values: [q for line in values for q in line])
        result = result.groupby(keys, dropna=False, as_index=False).agg(aggregations)
    result["is_document_anomaly"] = False
    result["document_anomaly_reason"] = ""
    result["document_threshold"] = np.nan
    eligible = result.quantity.gt(0) & result.document_type.eq("Расходная накладная")
    groups = result.loc[eligible].groupby(["brand", "sku"], sort=False)
    for _, group in groups:
        if len(group) < 8:
            continue
        q = group.quantity
        median = q.median()
        sigma = 1.4826 * (q - median).abs().median()
        if sigma == 0:
            sigma = (q.quantile(.75) - q.quantile(.25)) / 1.349
        threshold = max(median + 6 * sigma, 5 * median, 10.0)
        result.loc[group.index, "document_threshold"] = threshold
        for index in group.index[q > threshold]:
            if int((q >= q.loc[index] * .8).sum()) == 1:
                result.loc[index, ["is_document_anomaly", "document_anomaly_reason"]] = [True, "isolated_large_document"]
            else:
                result.loc[index, "document_anomaly_reason"] = "repeated_large_documents_not_excluded"
    return result


def reconcile_document_sales(history, documents):
    """Subtract document spikes only when shipment net equals the monthly total.

    A monthly blank is filled only by an observed shipment aggregate, never zero.
    Negative net returns stay auditable and are excluded from demand fitting.
    """
    shipments = documents.loc[documents.document_type.eq("Расходная накладная")].copy()
    shipments["date"] = shipments.date.dt.to_period("M").dt.to_timestamp()
    shipments["excluded"] = shipments.quantity.where(shipments.is_document_anomaly, 0)
    aggregate = shipments.groupby(["sku", "date"]).agg(
        document_net=("quantity", lambda s: s.sum(min_count=1)),
        document_missing_quantities=("quantity", lambda s: int(s.isna().sum())),
        document_excluded_qty=("excluded", "sum"), document_spikes=("is_document_anomaly", "sum"),
    ).reset_index()
    aggregate["observed_document_net"] = aggregate.document_net
    # A sum of only known lines is not a complete monthly total.
    aggregate.loc[aggregate.document_missing_quantities.gt(0), "document_net"] = np.nan
    result = history.merge(aggregate, on=["sku", "date"], how="left", validate="one_to_one")
    result["sales_source"] = np.where(result.raw_net_sales.notna(), "monthly_report", "unavailable")
    fill = result.raw_net_sales.isna() & result.document_net.notna()
    result.loc[fill, "raw_net_sales"] = result.loc[fill, "document_net"]
    result.loc[fill, "sales_source"] = "observed_document_net"
    result["document_reconciled"] = (np.isclose(result.raw_net_sales, result.document_net, atol=1e-6, rtol=0)
                                     & result.raw_net_sales.ge(result.document_excluded_qty.fillna(0)))
    result["document_anomalies_unreconciled"] = result.document_spikes.fillna(0).where(~result.document_reconciled, 0)
    result["document_anomalies_detected"] = result.document_spikes.fillna(0).where(result.document_reconciled, 0)
    result["document_excluded_qty"] = result.document_excluded_qty.fillna(0).where(result.document_reconciled, 0)
    result["raw_sales"] = result.raw_net_sales.where(result.raw_net_sales >= 0)
    result["sales"] = (result.raw_sales - result.document_excluded_qty).clip(lower=0)
    result["negative_net_returns"] = result.raw_net_sales.lt(0)
    return result


def load_partner_data(raw_dir) -> PartnerData:
    root = Path(raw_dir)
    books, inspection, issues = {}, [], []
    for filename in [name for files in FILES.values() for name in files.values()]:
        path = root / filename
        if not path.is_file():
            raise ValueError(f"Required partner workbook is missing: {filename}")
        wb = load_workbook(path, read_only=True, data_only=True)
        books[filename] = {}
        for sheet in wb:
            rows = list(sheet.values)
            books[filename][sheet.title] = rows
            inspection.append({"filename": filename, "sheet": sheet.title,
                               "rows": sheet.max_row, "columns": sheet.max_column})
        wb.close()
    catalogs, histories, doc_frames, inbound_frames, warehouse_frames, profiles = [], [], [], [], [], {}
    as_of = pd.Timestamp("2026-09-22")  # Explicit dated snapshot supplied by both transit reports.
    for brand, files in FILES.items():
        def table(role, required, header=0):
            rows = next(iter(books[files[role]].values()))
            return _table(rows, header, required, files[role])
        sales = table("sales", ["Номенклатура.Код", "Номенклатура"])
        inventory = table("inventory", ["Номенклатура.Код", "Номенклатура"])
        monthly = normalize_monthly(sales, "Номенклатура.Код", "Номенклатура", "raw_net_sales")
        stocks = normalize_monthly(inventory, "Номенклатура.Код", "Номенклатура", "historical_stock")
        doc = table("documents", ["Дата", "Номер", "Документ", "Код", "Номенклатура", "Склад", "Количество"])
        doc = _key_rows(doc, "Код")
        doc["date"] = pd.to_datetime(doc["Дата"], format="%d.%m.%Y %H:%M:%S", errors="coerce")
        doc["quantity"] = pd.to_numeric(doc["Количество"], errors="coerce")
        doc["document_type"] = doc["Документ"].str.replace(r"\s+\d.*", "", regex=True)
        doc["document_id"] = doc["Номер"].astype(str) + "@" + doc.date.dt.strftime("%Y-%m-%d")
        doc["brand"] = brand
        doc = doc.rename(columns={"Номенклатура": "product_name", "Склад": "warehouse", "Ед.": "unit"})
        issues.append({"brand": brand, "issue": "document_missing_quantity", "count": int(doc.quantity.isna().sum())})
        documents = doc[["brand", "sku", "product_name", "date", "document_id", "document_type", "warehouse", "unit", "quantity"]].copy()
        documents["source_file"] = files["documents"]
        documents = detect_document_anomalies(documents)
        doc_frames.append(documents)
        moq = table("moq", ["Код 1с", "Мин. разр. к отгр."] if brand == "IEK" else ["Номенклатура.Код", "Кратность"])
        moq = _key_rows(moq, "Код 1с" if brand == "IEK" else "Номенклатура.Код")
        moq_col = "Мин. разр. к отгр." if brand == "IEK" else "Кратность"
        moq["constraint"] = pd.to_numeric(moq[moq_col], errors="coerce")
        conflicting = moq.groupby("sku").constraint.nunique().gt(1)
        if conflicting.any():
            issues.append({"brand": brand, "issue": "conflicting_moq_unavailable", "count": int(conflicting.sum())})
            moq.loc[moq.sku.isin(conflicting.index[conflicting]), "constraint"] = np.nan
        moq = moq.drop_duplicates("sku")
        transit = table("transit", ["Код 1с", "Артикул ИЭК"] if brand == "IEK" else ["Код 1с", "Свободный остаток", "СЭ в пути 24.09"], header=0 if brand == "IEK" else 1)
        transit = _key_rows(transit, "Код 1с")
        if brand == "IEK":
            shipment_columns = [c for c in transit if "поступление до" in c]
            if len(shipment_columns) != 6:
                raise ValueError("IEK transit shipment columns changed; inspect the workbook")
            inbound = transit.melt(id_vars=["sku"], value_vars=shipment_columns, var_name="shipment", value_name="quantity")
            inbound["quantity"] = pd.to_numeric(inbound.quantity, errors="coerce")
            inbound = inbound.dropna(subset=["quantity"])
            inbound["arrival_date"] = pd.to_datetime(inbound.shipment.str.extract(r"поступление до (\d{2}\.\d{2}\.\d{4})")[0], format="%d.%m.%Y")
            inbound = inbound.drop_duplicates(["sku", "shipment", "quantity"])
        else:
            inbound = transit[["sku", "СЭ в пути 24.09"]].rename(columns={"СЭ в пути 24.09": "quantity"})
            inbound["shipment"] = "СЭ в пути 24.09"
            inbound["arrival_date"] = pd.Timestamp("2026-09-24")
            for warehouse in ["Витрина", "Остаток ТЗ", "РЦ ЕКТ  Рыскулова", "Розничный склад"]:
                part = transit[["sku", warehouse]].rename(columns={warehouse: "stock"})
                part["warehouse"] = warehouse
                part["brand"] = brand
                warehouse_frames.append(part)
        inbound["brand"] = brand
        inbound["source_file"] = files["transit"]
        inbound_frames.append(inbound)
        # Catalog union preserves products without usable sales rather than inventing history.
        names = [monthly[["sku", "product_name"]], stocks[["sku", "product_name"]], documents[["sku", "product_name"]],
                 transit[["sku", "Наименование"]].rename(columns={"Наименование": "product_name"}),
                 moq[["sku", "Наименование" if brand == "IEK" else "Номенклатура"]].rename(columns={"Наименование": "product_name", "Номенклатура": "product_name"})]
        names = pd.concat(names)
        # A blank name in the first report must not mask a supplied name elsewhere.
        names["product_name"] = names.product_name.replace(r"^\s*$", np.nan, regex=True)
        catalog = names.sort_values("product_name", key=lambda s: s.isna(), kind="stable").drop_duplicates("sku").set_index("sku")
        catalog["brand"] = catalog["supplier"] = brand
        catalog["warehouse"] = None  # Not allocated: document warehouse is not inventory scope.
        catalog["inventory_scope"] = "all_reported_warehouses"
        catalog["category"] = None
        catalog["lead_time_days"] = np.nan
        catalog["lead_time_source"] = "unavailable"
        catalog["unit_price"] = np.nan  # СС реал is cost, not confirmed supplier price.
        catalog["minimum_order_qty"] = moq.set_index("sku").constraint if brand == "IEK" else np.nan
        catalog["order_multiple"] = moq.set_index("sku").constraint if brand != "IEK" else np.nan
        latest_month = stocks.date.max()
        catalog["current_stock"] = stocks.loc[stocks.date.eq(latest_month)].set_index("sku").historical_stock
        catalog["stock_as_of"] = latest_month
        catalog["stock_source"] = "latest_monthly_inventory"
        catalog["in_transit"] = inbound.groupby("sku").quantity.sum(min_count=1)
        catalog["transit_source"] = np.where(catalog.in_transit.notna(), "provided_shipment_quantities", "unavailable")
        if brand != "IEK":
            snap = transit.set_index("sku")
            catalog.loc[snap.index, "current_stock"] = snap["Свободный остаток"]
            catalog.loc[snap.index, "stock_as_of"] = as_of
            catalog.loc[snap.index, "stock_source"] = "provided_free_stock_snapshot"
            catalog.loc[snap.index, "category"] = snap["Категория 2026"]
            catalog["reported_cost"] = snap["СС реал"]
        months = pd.date_range(monthly.date.min(), monthly.date.max(), freq="MS")
        hist = pd.MultiIndex.from_product([catalog.index, months], names=["sku", "date"]).to_frame(index=False)
        hist = hist.merge(monthly[["sku", "date", "raw_net_sales"]], on=["sku", "date"], how="left", validate="one_to_one")
        hist = hist.merge(stocks[["sku", "date", "historical_stock"]], on=["sku", "date"], how="left", validate="one_to_one")
        hist = reconcile_document_sales(hist, documents)
        hist = infer_stockout_evidence(hist)
        hist["partial_month"] = hist.date.dt.to_period("M").eq(as_of.to_period("M"))
        hist["brand"] = brand
        catalogs.append(catalog.reset_index())
        histories.append(hist)
        season_rows = next(iter(books[files["seasonality"]].values()))
        if str(season_rows[9][11]).strip() != "СЕЗОННОСТЬ":
            raise ValueError(f"{files['seasonality']}: expected supplied seasonality coefficient column")
        profile = {MONTHS[str(r[1]).strip()[:3]]: float(r[11]) for r in season_rows[10:22]}
        if len(profile) != 12 or not all(np.isfinite(v) and v > 0 for v in profile.values()):
            issues.append({"brand": brand, "issue": "unusable_partner_seasonality", "count": 1})
        else:
            profiles[brand] = profile
    catalog = pd.concat(catalogs, ignore_index=True)
    if catalog.sku.duplicated().any():
        raise ValueError("1C code collision across brands requires explicit disambiguation")
    return PartnerData(catalog, pd.concat(histories, ignore_index=True), pd.concat(doc_frames, ignore_index=True),
                       pd.concat(inbound_frames, ignore_index=True), profiles,
                       pd.concat(warehouse_frames, ignore_index=True), inspection, issues, as_of)
